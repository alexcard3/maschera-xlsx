"""Genera un Excel di prova con dati INVENTATI per collaudare l'anonimizzatore.

Uso: python genera_prova.py <cartella di uscita>
Scrive prova_bilancio.xlsx e sensibili.json (i valori che non devono restare in chiaro), prova_prospetto.xlsx e
prospetto.json (nomi sparsi in un prospetto).
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment

DISPARI = [1, 0, 5, 7, 9, 13, 15, 17, 19, 21, 2, 4, 18, 20, 11, 3, 6, 8, 12, 14, 16, 10, 22, 25, 24, 23]


def cf_controllo(primi15):
    s = 0
    for i, ch in enumerate(primi15):
        v = int(ch) if ch.isdigit() else ord(ch) - 65
        s += DISPARI[v] if i % 2 == 0 else v
    return chr(65 + s % 26)


def luhn_ok(numero):
    s = 0
    for i, d in enumerate(reversed(numero)):
        x = int(d)
        if i % 2 == 1:
            x *= 2
            if x > 9:
                x -= 9
        s += x
    return s % 10 == 0


def piva(primi10):
    for c in "0123456789":
        if luhn_ok(primi10 + c):
            return primi10 + c


def main(uscita):
    uscita = Path(uscita)
    uscita.mkdir(parents=True, exist_ok=True)
    assert cf_controllo("RSSMRA85T10A562") == "S"   # esempio pubblico noto
    cf1 = "VRDLRA79C45L219" + cf_controllo("VRDLRA79C45L219")
    cf2 = "FBBGNN66M12F205" + cf_controllo("FBBGNN66M12F205")
    pive = [piva(p) for p in ("0432176891", "1187654320", "0299887766", "0712345098", "0155566677", "0900011122")]
    piva_numero = piva("0876543210")
    iban = "IT60X0542811101000000123456"   # esempio pubblico noto, mod-97 valido
    # importo intero di 10 cifre che per caso passa il controllo della P.IVA: deve restare un numero
    importo_come_piva = next(n for n in range(1234567890, 1234567990) if luhn_ok(str(n).zfill(11)))
    # nome con l'accento separato dalla lettera (NFD), come nei testi incollati da un Mac
    nfd = "Nicolo\u0300 Ferre\u0300"

    clienti = [
        ("Edilnord Srl", "Marco Galli", "m.galli@edilnord.it", "011 5551234", pive[0], 152340.50,
         "Cliente storico, collegato a Tecnofarm S.p.A."),
        ("Tecnofarm S.p.A.", "Giovanna Ferrero", "amministrazione@tecnofarm.com", "+39 347 1234567", pive[1], 98760.00,
         "Sollecito inviato a Gianni Fabbri il 03/02/2025"),
        ("Bianchi & Figli Snc", "Paolo Bianchi", "paolo@bianchiefigli.it", "02 48001234", pive[2], 45000.00, ""),
        ("Studio Associato Rinaldi", "Elena Rinaldi", "elena.rinaldi@studiorinaldi.it", "333 9876543", pive[3], 12500.75,
         "Referente sostituita da Elena Rinaldi a marzo"),
        ("Alfa Logistica Srl", "Davide Conti", "d.conti@alfalogistica.it", "0125 612345", pive[4], 230000.00,
         "Contratto quadro firmato con Davide Conti"),
        ("Mondo Verde Coop", "Sara Moretti", "info@mondoverde.coop", "011 7778899", pive[5], 8700.00, "Verificare P.IVA"),
    ]
    movimenti = [
        ("2025-01-15", "Fattura Edilnord Srl gennaio", 12500.00, "CM-2291", "Edilnord Srl"),
        ("2025-02-03", f"Bonifico a Mario Bianchi IBAN {iban} rata 2", -3200.00, "CM-2291", "Bianchi & Figli Snc"),
        ("2025-02-20", f"Rimborso spese Laura Verdi CF {cf1}", -145.30, "AMM-01", "Laura Verdi"),
        ("2025-03-05", "Pagamento Tecnofarm S.p.A. tel 011 5551234", 8900.00, "TF-778", "Tecnofarm S.p.A."),
        ("2025-03-18", f"Nota di credito, contatto Gianni Fabbri ({cf2})", -900.00, "TF-778", "Tecnofarm S.p.A."),
        ("2025-04-01", "Saldo commessa CM-2291 del 01/04/2025, fattura n. 2025/0012", 4300.00, "CM-2291", "Alfa Logistica Srl"),
        ("2025-04-30", "Accantonamento rischi su crediti", -5000.00, "", ""),
        ("2025-05-10", "Incassi cumulati in centesimi", importo_come_piva, "", ""),
        ("2025-05-12", "Telefonata con " + nfd, 0.00, "", ""),
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Clienti"
    ws["A1"] = "Bilancio clienti 2025 - preparato da Laura Verdi"
    ws.append([])
    ws.append(["Ragione sociale", "Referente", "Email", "Telefono", "P.IVA", "Fatturato 2025", "Note"])
    for r in clienti:
        ws.append(list(r))
    ws["F4"].comment = Comment("Dato confermato da Marco Galli", "Laura Verdi")

    ws2 = wb.create_sheet("Movimenti")
    ws2.append(["Data", "Descrizione", "Importo", "Commessa", "Controparte"])
    from datetime import date
    for d, desc, imp, comm, contr in movimenti:
        ws2.append([date.fromisoformat(d), desc, imp, comm, contr or None])
    ws2["G1"] = "Totale"
    ws2["G2"] = "=SUM(C2:C8)"

    ws3 = wb.create_sheet("Note Edilnord")
    ws3.sheet_state = "hidden"
    ws3["A1"] = "Riunione con Marco Galli e Sara Moretti"
    ws3["A2"] = "Scrivere a"
    ws3["B2"] = "m.galli@edilnord.it"
    ws3["B2"].hyperlink = "mailto:m.galli@edilnord.it"
    ws3["A3"] = "P.IVA salvata come numero"
    ws3["B3"] = int(pive[1])
    # P.IVA che nel file c'e' solo come numero, con lo zero iniziale perso: al ripristino torna un numero
    ws3["A4"] = "Partita IVA fornitore"
    ws3["B4"] = int(piva_numero)

    wb.properties.creator = "Laura Verdi"
    wb.save(uscita / "prova_bilancio.xlsx")

    sensibili = sorted({
        *(c[0] for c in clienti), *(c[1] for c in clienti), *(c[2] for c in clienti), *(c[3] for c in clienti),
        *pive, piva_numero, piva_numero.lstrip("0"), iban, cf1, cf2, "Laura Verdi", "Gianni Fabbri", "Mario Bianchi", "CM-2291", "Edilnord",
        "Nicol\u00f2 Ferr\u00e8",
    })
    numeri = {"Clienti!F": [c[5] for c in clienti], "Movimenti!C": [m[2] for m in movimenti]}
    with open(uscita / "sensibili.json", "w", encoding="utf-8", newline="") as fh:
        json.dump({"sensibili": sensibili, "numeri": numeri, "pive": pive, "cf": [cf1, cf2], "iban": iban,
                   "importo_come_piva": importo_come_piva, "piva_numero": piva_numero},
                  fh, ensure_ascii=False, indent=1)
    prospetto(uscita)
    print("ok", uscita)


def prospetto(uscita):
    """Prospetto a blocchi senza intestazioni: unita' e cognome in testa al blocco, mesi e voci sotto."""
    cognomi = ["Ferrandi", "Lupetti", "Casoli", "Merloni", "Bertagna", "Zoccali",
               "Pirovano", "Tessaro", "Valsecchi", "DallaRosa", "Morandini", "Ferrandi"]   # Ferrandi due volte
    unita = ["SB1 T", "2T", "3 T", "4 T", "5 T", "6 T", "7 T", "8 T", "9 T", "1 C", "2 C", "3 C"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Consumi"
    numeri = []
    riga = 1
    for b in range(4):
        for k in range(3):
            i = b * 3 + k
            col = 1 + 3 * k
            ws.cell(riga, col, unita[i])
            ws.cell(riga, col + 1, cognomi[i])
            valori = [300 + 10 * i, 5 + i, 20.5 + i]
            for j, (voce, v) in enumerate(zip(["Gennaio", "LUCE", "Risc"], valori)):
                ws.cell(riga + 2 + j, col, voce)
                ws.cell(riga + 2 + j, col + 2, v)
                numeri.append([riga + 2 + j, col + 2, v])
            ws.cell(riga + 6, col, "TOTALE")
            ws.cell(riga + 6, col + 2, sum(valori))
            numeri.append([riga + 6, col + 2, sum(valori)])
        riga += 8
    wb.save(uscita / "prova_prospetto.xlsx")
    with open(uscita / "prospetto.json", "w", encoding="utf-8", newline="") as fh:
        json.dump({"cognomi": sorted(set(cognomi)), "numeri": numeri,
                   "etichette": ["Gennaio", "LUCE", "Risc", "TOTALE", "SB1 T", "2T"]}, fh, ensure_ascii=False, indent=1)


if __name__ == "__main__":
    main(sys.argv[1])
