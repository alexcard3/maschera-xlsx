"""Banco di prova con dati INVENTATI: quanto aiuta rizzo-pii a proporre i nomi in un foglio di calcolo.

Scrive nella cartella indicata:
- banco.xlsx   quattro fogli (fornitori, movimenti, personale, appunti), creato con openpyxl: nessun autore
- celle.json   le celle di testo con foglio, indirizzo, intestazione di colonna e testo della riga
I nomi da trovare (la verita') stanno in NOMI, con il gruppo a cui appartengono; valuta.py ne cerca le occorrenze.
Uso: python banco.py <cartella>
"""
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

FOGLI = {
    "Fornitori": [
        ["Ragione sociale", "Referente", "Città", "Importo", "Note"],
        ["Edilnord Srl", "Marco Galli", "Torino", 152340.5, "Contratto rinnovato a marzo"],
        ["Tecnofarm S.p.A.", "Giovanna Ferrero", "Milano", 98760, "Sollecito inviato a Gianni Fabbri"],
        ["Bianchi & Figli Snc", "PAOLO BIANCHI", "Asti", 45000, None],
        ["Studio Associato Rinaldi", "elena rinaldi", "Cuneo", 12500.75, "Referente sostituita a giugno"],
        ["Alfa Logistica Srl", "Conti Davide", "Novara", 230000, "Pagamento a 60 giorni"],
        ["Mondo Verde Coop", "Dott.ssa Sara Moretti", "Biella", 8700, "Verificare P.IVA"],
        ["F.lli Ruggeri", "Ruggeri", "Alba", 3100, "Ordine urgente di Luca Ruggeri"],
        ["Officina Meccanica Pautasso", "G. Pautasso", "Pinerolo", 5600, None],
        ["OMNIATECH", "ANNA COLOMBO", "Ivrea", 19800, "Nuovo fornitore"],
        ["Grafiche Serra di Fabio Serra", "fabio serra", "Bra", 2300, "Fattura contestata da Marta Gentile"],
        ["Caseificio Valle Stura", "Bruno Giordano", "Mondovì", 7400, "Consegna ogni lunedì"],
        ["Lumen Energia S.r.l.", "Ing. Chiara Fogliato", "Torino", 61000, "Contratto quadro firmato con Chiara Fogliato"],
    ],
    "Movimenti": [
        ["Data", "Descrizione", "Controparte", "Importo", "Causale"],
        [date(2025, 1, 15), "Fattura Edilnord gennaio", "Edilnord Srl", 12500, "Acquisti"],
        [date(2025, 2, 3), "Bonifico a Marco Galli rimborso spese", "Marco Galli", -320, "Rimborsi"],
        [date(2025, 2, 28), "Stipendio febbraio", None, -2100, "Personale"],
        [date(2025, 3, 10), "Incasso da Tecnofarm", "TECNOFARM S.P.A.", 9876, "Vendite"],
        [date(2025, 3, 31), "Accantonamento fondo rischi", None, -5000, "Accantonamenti"],
        [date(2025, 4, 12), "Nota di credito a favore di Bianchi & Figli", "Bianchi & Figli Snc", -450, "Resi"],
        [date(2025, 5, 5), "Anticipo a Davide Conti per trasferta Roma", "conti davide", -600, "Trasferte"],
        [date(2025, 6, 30), "Saldo iniziale", None, 0, "Apertura"],
        [date(2025, 7, 1), "Canone leasing auto aziendale", "Lumen Energia S.r.l.", -890, "Leasing"],
        [date(2025, 7, 18), "Pagamento fattura n. 2025/118", "OMNIATECH", -19800, "Acquisti"],
        [date(2025, 9, 2), "Quota associativa Confartigianato", "Confartigianato", -250, "Quote"],
        [date(2025, 10, 1), "Totale terzo trimestre", None, None, "Totali"],
    ],
    "Personale": [
        ["COGNOME", "NOME", "REPARTO", "COSTO ANNUO", "NOTE"],
        ["ROSSI", "MARIO", "AMMINISTRAZIONE", 42000, "PART TIME"],
        ["FERRARI", "GIULIA", "VENDITE", 38000, None],
        ["ESPOSITO", "LUCA", "MAGAZZINO", 31000, "TRASFERITO DA NAPOLI"],
        ["RICCI", "FRANCESCA", "ACQUISTI", 35000, None],
        ["MARINO", "ANDREA", "PRODUZIONE", 29500, "IN MALATTIA DAL 12/03"],
        ["GRECO", "SOFIA", "PRODUZIONE", 29500, None],
    ],
    "Appunti": [
        ["Appunti"],
        ["Riunione con Marco Galli e Sara Moretti del 12 marzo: rinvio del pagamento."],
        ["Il dott. Bruno Giordano chiede copia del contratto."],
        ["Chiamare Anna Colombo (Omniatech) entro venerdì."],
        ["Budget 2026 approvato dal CdA."],
        ["Verificare con il commercialista lo scadenzario."],
    ],
}

# La verita': ogni forma scritta di un nome, con il suo gruppo. Si cerca tale e quale (maiuscole comprese).
NOMI = {
    "persona, iniziali maiuscole": ["Marco Galli", "Giovanna Ferrero", "Gianni Fabbri", "Sara Moretti", "Luca Ruggeri",
                                    "Marta Gentile", "Bruno Giordano", "Chiara Fogliato", "Anna Colombo", "Davide Conti"],
    "persona, tutto maiuscolo": ["PAOLO BIANCHI", "ANNA COLOMBO"],
    "persona, tutto minuscolo": ["elena rinaldi", "fabio serra", "conti davide"],
    "persona, cognome prima o iniziale": ["Conti Davide", "G. Pautasso"],
    "persona, solo nome o solo cognome": ["Ruggeri", "ROSSI", "MARIO", "FERRARI", "GIULIA", "ESPOSITO", "LUCA", "RICCI",
                                          "FRANCESCA", "MARINO", "ANDREA", "GRECO", "SOFIA"],
    "azienda, con forma societaria": ["Edilnord Srl", "Tecnofarm S.p.A.", "Bianchi & Figli Snc", "Alfa Logistica Srl",
                                      "Mondo Verde Coop", "F.lli Ruggeri", "Lumen Energia S.r.l."],
    "azienda, senza forma societaria": ["Studio Associato Rinaldi", "Officina Meccanica Pautasso",
                                        "Grafiche Serra di Fabio Serra", "Caseificio Valle Stura", "Confartigianato",
                                        "Edilnord", "Tecnofarm", "Bianchi & Figli", "Omniatech"],
    "azienda, tutto maiuscolo": ["OMNIATECH", "TECNOFARM S.P.A."],
}


def main(cartella):
    cartella = Path(cartella)
    cartella.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    celle = []
    for nome, righe in FOGLI.items():
        ws = wb.create_sheet(nome)
        intest = righe[0]
        for r, riga in enumerate(righe, start=1):
            ws.append(riga)
            testo_riga = "; ".join("%s: %s" % (intest[c], v) for c, v in enumerate(riga)
                                   if isinstance(v, str) and r > 1) if r > 1 else ""
            for c, v in enumerate(riga):
                if isinstance(v, str):
                    celle.append({"foglio": nome, "cella": "%s%d" % (get_column_letter(c + 1), r), "riga": r,
                                  "intestazione": intest[c] if r > 1 else None, "testo": v, "testo_riga": testo_riga})
    wb.properties.creator = None
    wb.properties.lastModifiedBy = None
    wb.save(cartella / "banco.xlsx")
    with open(cartella / "celle.json", "w", encoding="utf-8") as fh:
        json.dump({"celle": celle, "nomi": NOMI}, fh, ensure_ascii=False, indent=1)
    print("%d celle di testo, %d forme di nome" % (len(celle), sum(len(v) for v in NOMI.values())))


if __name__ == "__main__":
    main(sys.argv[1])
